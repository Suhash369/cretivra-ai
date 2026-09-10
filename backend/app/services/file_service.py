import os
import shutil
import uuid
import base64
from typing import Dict, Any, Optional
import pypdf
import docx
from PIL import Image
import pptx

from app.core.config import settings
from app.core.security import sanitize_filename, validate_path_safety
from app.core.logging import logger

ALLOWED_EXTENSIONS = {
    ".pdf", ".docx", ".pptx", ".ppt", ".txt", ".csv", ".md",
    ".png", ".jpg", ".jpeg", ".webp"
}

ALLOWED_MIME_TYPES = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
    "application/vnd.ms-powerpoint": ".ppt",
    "text/plain": ".txt",
    "text/csv": ".csv",
    "text/markdown": ".md",
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/webp": ".webp"
}

class FileService:
    def __init__(self, upload_dir: str = settings.UPLOAD_DIR):
        self.upload_dir = os.path.abspath(upload_dir)
        os.makedirs(self.upload_dir, exist_ok=True)

    def validate_file(self, filename: str, file_size: int) -> Dict[str, Any]:
        ext = os.path.splitext(filename)[1].lower()
        if ext not in ALLOWED_EXTENSIONS:
            return {
                "valid": False,
                "error": f"Unsupported file extension '{ext}'. Allowed extensions: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
            }
        
        max_bytes = settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024
        if file_size > max_bytes:
            return {
                "valid": False,
                "error": f"File size exceeds maximum allowed size of {settings.MAX_UPLOAD_SIZE_MB}MB."
            }

        return {"valid": True}

    async def save_file(self, file_bytes: bytes, filename: str, mime_type: str) -> Dict[str, Any]:
        validation = self.validate_file(filename, len(file_bytes))
        if not validation["valid"]:
            raise ValueError(validation["error"])

        clean_name = sanitize_filename(filename)
        file_id = str(uuid.uuid4())
        stored_filename = f"{file_id}_{clean_name}"
        file_path = os.path.join(self.upload_dir, stored_filename)

        if not validate_path_safety(file_path, self.upload_dir):
            raise ValueError("Invalid file path / potential path traversal detected.")

        with open(file_path, "wb") as f:
            f.write(file_bytes)

        ext = os.path.splitext(clean_name)[1].lower()
        data_url = None
        image_metadata = None

        # Process image data if image
        if ext in [".png", ".jpg", ".jpeg", ".webp"]:
            try:
                with Image.open(file_path) as img:
                    image_metadata = {
                        "width": img.width,
                        "height": img.height,
                        "format": img.format,
                        "mode": img.mode
                    }
                b64 = base64.b64encode(file_bytes).decode("utf-8")
                safe_mime = mime_type if mime_type and mime_type.startswith("image/") else f"image/{ext.replace('.', '')}"
                data_url = f"data:{safe_mime};base64,{b64}"
            except Exception as img_err:
                logger.warning(f"Failed to read image metadata for {clean_name}: {img_err}")

        # Extract text content if document or image
        extracted_text = self.extract_text_content(file_path, clean_name, image_metadata=image_metadata)

        return {
            "id": file_id,
            "filename": clean_name,
            "mime_type": mime_type,
            "path": file_path,
            "size": len(file_bytes),
            "extracted_text": extracted_text,
            "data_url": data_url,
            "image_metadata": image_metadata
        }

    def extract_text_content(
        self,
        file_path: str,
        filename: str,
        image_metadata: Optional[Dict[str, Any]] = None
    ) -> Optional[str]:
        ext = os.path.splitext(filename)[1].lower()
        try:
            if ext in [".txt", ".md", ".csv"]:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()

            elif ext == ".pdf":
                reader = pypdf.PdfReader(file_path)
                text_pages = []
                for idx, page in enumerate(reader.pages, 1):
                    t = page.extract_text()
                    if t and t.strip():
                        text_pages.append(f"--- Page {idx} ---\n{t.strip()}")
                return "\n\n".join(text_pages) if text_pages else "Empty or scanned PDF document."

            elif ext == ".docx":
                doc = docx.Document(file_path)
                paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                return "\n\n".join(paras) if paras else "Empty Word document."

            elif ext in [".pptx", ".ppt"]:
                prs = pptx.Presentation(file_path)
                slides_output = []
                for idx, slide in enumerate(prs.slides, 1):
                    slide_lines = [f"=== Slide {idx} ==="]
                    
                    # Extract shape text and tables
                    for shape in slide.shapes:
                        if shape.has_text_frame and shape.text.strip():
                            slide_lines.append(shape.text.strip())
                        elif shape.has_table:
                            table_rows = []
                            for row in shape.table.rows:
                                cells_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                                if cells_text:
                                    table_rows.append(" | ".join(cells_text))
                            if table_rows:
                                slide_lines.append("\n".join(table_rows))
                    
                    # Extract speaker notes if present
                    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                        notes = slide.notes_slide.notes_text_frame.text.strip()
                        if notes:
                            slide_lines.append(f"[Speaker Notes]: {notes}")
                    
                    slides_output.append("\n".join(slide_lines))
                
                return "\n\n".join(slides_output) if slides_output else "Empty PowerPoint presentation."

            elif ext in [".png", ".jpg", ".jpeg", ".webp"]:
                dims = f"{image_metadata['width']}x{image_metadata['height']}" if image_metadata else "Unknown"
                fmt = image_metadata.get("format", ext.replace(".", "").upper()) if image_metadata else ext.upper()
                return (
                    f"[Attached Image File]: {filename}\n"
                    f"Format: {fmt} | Resolution: {dims}\n"
                    f"This image has been attached by the user. Please examine and answer questions about it directly."
                )

        except Exception as e:
            logger.error(f"Error extracting text from {filename}: {e}")
        return None

    async def describe_image_with_vision(
        self,
        image_bytes: bytes,
        mime_type: str = "image/jpeg",
        prompt: str = "Thoroughly inspect and explain the visual contents, text, UI elements, and data in this image."
    ) -> Optional[str]:
        """
        Multimodal visual analysis helper using Google Gemini Vision or OpenRouter Vision.
        """
        b64_data = base64.b64encode(image_bytes).decode("utf-8")
        safe_mime = mime_type if mime_type.startswith("image/") else "image/jpeg"

        # 1. Try Gemini Vision if key exists
        gemini_key = getattr(settings, "GEMINI_API_KEY", "")
        if gemini_key:
            clean_key = re.sub(r'[\r\n\t ]+', '', gemini_key)
            for model_name in ["gemini-3.6-flash", "gemini-3.7-flash", "gemini-flash-latest"]:
                try:
                    import httpx
                    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent?key={clean_key}"
                    payload = {
                        "contents": [{
                            "role": "user",
                            "parts": [
                                {"text": prompt},
                                {"inline_data": {"mime_type": safe_mime, "data": b64_data}}
                            ]
                        }]
                    }
                    async with httpx.AsyncClient(timeout=30.0) as client:
                        res = await client.post(url, json=payload)
                        if res.status_code == 200:
                            data = res.json()
                            parts = data.get("candidates", [{}])[0].get("content", {}).get("parts", [])
                            text_out = "".join([p.get("text", "") for p in parts if "text" in p])
                            if text_out:
                                return text_out.strip()
                except Exception as g_err:
                    logger.warning(f"Gemini describe image error: {g_err}")

        # 2. Try OpenRouter Vision if key exists
        openrouter_key = getattr(settings, "OPENROUTER_API_KEY", "")
        if openrouter_key:
            try:
                import httpx
                data_url = f"data:{safe_mime};base64,{b64_data}"
                payload = {
                    "model": "inclusionai/ling-3.0-flash-vl:free",
                    "messages": [{
                        "role": "user",
                        "content": [
                            {"type": "text", "text": prompt},
                            {"type": "image_url", "image_url": {"url": data_url}}
                        ]
                    }]
                }
                headers = {
                    "Authorization": f"Bearer {openrouter_key}",
                    "Content-Type": "application/json"
                }
                async with httpx.AsyncClient(timeout=30.0) as client:
                    res = await client.post("https://openrouter.ai/api/v1/chat/completions", headers=headers, json=payload)
                    if res.status_code == 200:
                        content = res.json().get("choices", [{}])[0].get("message", {}).get("content", "")
                        if content:
                            return content.strip()
            except Exception as or_err:
                logger.warning(f"OpenRouter describe image error: {or_err}")

        return None

file_service = FileService()
